import datetime
import html as html_lib
import re
import sqlite3
import sys
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
import requests


# Windows terminals may default to a legacy code page that cannot print the
# Vietnamese/Unicode status messages used by this crawler.
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")


def parse_mega645_html(html_text):
    """Parse the latest Mega 6/45 result from the HTML page when the API is unavailable."""
    if not html_text:
        return None

    # Official single-draw page. Parse this structured block before the
    # looser history-page expressions below so an exact requested draw ID is
    # never confused with a navigation or previous-result value in the page.
    official_draw = re.search(
        r"<h5>\s*[^<]*<b>\s*#?(\d+)\s*</b>\s*[^<]*<b>\s*(\d{1,2}/\d{1,2}/\d{4})\s*</b>",
        html_text,
        re.I | re.S,
    )
    official_numbers = re.search(
        r'<div[^>]*class=["\']day_so_ket_qua_v2["\'][^>]*>(.*?)</div>',
        html_text,
        re.I | re.S,
    )
    if official_draw and official_numbers:
        nums = [
            int(value)
            for value in re.findall(
                r'class=["\'][^"\']*bong_tron[^"\']*["\'][^>]*>\s*(\d{1,2})\s*<',
                official_numbers.group(1),
                re.I,
            )
        ]
        if len(nums) == 6 and len(set(nums)) == 6:
            draw_date_str = official_draw.group(2)
            dt = datetime.datetime.strptime(draw_date_str, "%d/%m/%Y")
            weekdays = [
                "Thứ 2", "Thứ 3", "Thứ 4", "Thứ 5", "Thứ 6", "Thứ 7", "Chủ Nhật",
            ]
            jackpot_match = re.search(
                r'class=["\']so_tien["\'][^>]*>\s*<h3>\s*([\d.,]+)\s*</h3>',
                html_text,
                re.I | re.S,
            )
            winner_match = re.search(
                r'<td>\s*Jackpot\s*</td>.*?<td[^>]*>\s*([\d.,]+)\s*</td>',
                html_text,
                re.I | re.S,
            )
            parse_amount = lambda value: int(value.replace(".", "").replace(",", "")) if value else 0
            return {
                "draw_id": f"#{int(official_draw.group(1)):05d}",
                "date": draw_date_str,
                "weekday": weekdays[dt.weekday()],
                "nums": sorted(nums),
                "jackpot": parse_amount(jackpot_match.group(1) if jackpot_match else ""),
                "winners": parse_amount(winner_match.group(1) if winner_match else ""),
            }

    raw_html = html_text
    text = html_lib.unescape(html_text)
    text = text.replace("&nbsp;", " ").replace("\xa0", " ")
    text = re.sub(r"<[^>]+>", " ", text, flags=re.S)
    text = re.sub(r"\s+", " ", text).strip()

    live_rows = re.finditer(
        r"<td[^>]*>\s*(\d{1,2}/\d{1,2}/\d{4})\s*</td>\s*"
        r"<td[^>]*>\s*(?:<a[^>]*>)?\s*#?0*(\d+)\s*(?:</a>)?\s*</td>\s*"
        r"<td[^>]*>(.*?)</td>",
        raw_html,
        re.I | re.S,
    )
    for match in live_rows:
        draw_date_str = match.group(1)
        draw_id = f"#{int(match.group(2)):05d}"
        cell_html = match.group(3)
        nums = [
            int(x)
            for x in re.findall(r"<span[^>]*>\s*(\d{1,2})\s*</span>", cell_html, re.I)
        ]
        if len(nums) == 6 and len(set(nums)) == 6:
            dt = datetime.datetime.strptime(draw_date_str, "%d/%m/%Y")
            weekdays = [
                "Thứ 2",
                "Thứ 3",
                "Thứ 4",
                "Thứ 5",
                "Thứ 6",
                "Thứ 7",
                "Chủ Nhật",
            ]
            weekday = weekdays[dt.weekday()]
            return {
                "draw_id": draw_id,
                "date": draw_date_str,
                "weekday": weekday,
                "nums": sorted(nums),
                "jackpot": 0,
                "winners": 0,
            }

    row_match = re.search(
        r"(\d{2}/\d{2}/\d{4})\s*\|\s*0*(\d+)\s*\|\s*(?P<nums>(?:\d{1,2}\s+){5}\d{1,2})",
        text,
        re.I,
    )
    if row_match:
        draw_date_str = row_match.group(1)
        draw_id = f"#{int(row_match.group(2)):05d}"
        nums = [int(x) for x in row_match.group("nums").split()]
        dt = datetime.datetime.strptime(draw_date_str, "%d/%m/%Y")
        weekdays = [
            "Thứ 2",
            "Thứ 3",
            "Thứ 4",
            "Thứ 5",
            "Thứ 6",
            "Thứ 7",
            "Chủ Nhật",
        ]
        weekday = weekdays[dt.weekday()]
        return {
            "draw_id": draw_id,
            "date": draw_date_str,
            "weekday": weekday,
            "nums": sorted(nums),
            "jackpot": 0,
            "winners": 0,
        }

    direct_match = re.search(
        r"(\d{2}/\d{2}/\d{4})\s*#\s*0*(\d+)\s*(?P<nums>(?:\d{1,2}\s+){5}\d{1,2})",
        text,
        re.I,
    )
    if direct_match:
        draw_date_str = direct_match.group(1)
        draw_id = f"#{int(direct_match.group(2)):05d}"
        nums = [int(x) for x in direct_match.group("nums").split()]
        dt = datetime.datetime.strptime(draw_date_str, "%d/%m/%Y")
        weekdays = [
            "Thứ 2",
            "Thứ 3",
            "Thứ 4",
            "Thứ 5",
            "Thứ 6",
            "Thứ 7",
            "Chủ Nhật",
        ]
        weekday = weekdays[dt.weekday()]
        return {
            "draw_id": draw_id,
            "date": draw_date_str,
            "weekday": weekday,
            "nums": sorted(nums),
            "jackpot": 0,
            "winners": 0,
        }

    match_draw = re.search(r"Kết quả QSMT kỳ\s*#?(\d+)\s*ngày\s*(\d{2}/\d{2}/\d{4})", text, re.I)
    if not match_draw:
        return None

    draw_id = f"#{int(match_draw.group(1)):05d}"
    draw_date_str = match_draw.group(2)
    dt = datetime.datetime.strptime(draw_date_str, "%d/%m/%Y")
    weekdays = [
        "Thứ 2",
        "Thứ 3",
        "Thứ 4",
        "Thứ 5",
        "Thứ 6",
        "Thứ 7",
        "Chủ Nhật",
    ]
    weekday = weekdays[dt.weekday()]

    num_block = re.search(
        r"ngày\s*\d{1,2}/\d{1,2}/\d{4}\s*(?P<nums>(?:\d{1,2}\s+){5}\d{1,2})",
        text,
        re.I | re.S,
    )
    if not num_block:
        return None

    nums = [int(x) for x in num_block.group("nums").split()]
    if len(nums) != 6 or len(set(nums)) != 6:
        return None

    jackpot_match = re.search(
        r"Jackpot\s+Mega\s+6/45\s+ước\s+tính.*?(?P<jackpot>[\d\.,]+)\s*VNĐ|(?P<jackpot2>[\d\.,]+)\s*VNĐ.*?Jackpot\s+Mega\s+6/45",
        text,
        re.I | re.S,
    )
    jackpot_val = 0
    if jackpot_match:
        jackpot_text = jackpot_match.group("jackpot") or jackpot_match.group("jackpot2")
        if jackpot_text:
            jackpot_val = int(jackpot_text.replace(".", "").replace(",", ""))

    return {
        "draw_id": draw_id,
        "date": draw_date_str,
        "weekday": weekday,
        "nums": sorted(nums),
        "jackpot": jackpot_val,
        "winners": 0,
    }


def parse_ketquadientoan_archive_html(html_text):
    """Parse the full Mega 6/45 history table from the archive page."""
    if not html_text:
        return []

    text = html_lib.unescape(html_text)
    rows = []
    for match in re.finditer(r"<tr[^>]*>(.*?)</tr>", text, re.I | re.S):
        block = match.group(1)
        if "home-mini-whiteball" not in block:
            continue

        date_match = re.search(r">(?:[A-Z]{1,2},\s*)?(\d{1,2}/\d{1,2}/\d{4})<", block, re.I)
        if not date_match:
            continue

        nums = [
            int(value)
            for value in re.findall(r'class=["\']home-mini-whiteball["\'][^>]*>(\d{1,2})<', block, re.I)
        ]
        if len(nums) != 6:
            continue

        draw_date_str = date_match.group(1)
        dt = datetime.datetime.strptime(draw_date_str, "%d/%m/%Y")
        weekdays = [
            "Thứ 2",
            "Thứ 3",
            "Thứ 4",
            "Thứ 5",
            "Thứ 6",
            "Thứ 7",
            "Chủ Nhật",
        ]
        rows.append({
            "draw_id": "",
            "date": draw_date_str,
            "weekday": weekdays[dt.weekday()],
            "nums": sorted(nums),
            "jackpot": 0,
            "winners": 0,
        })

    rows.sort(key=lambda row: datetime.datetime.strptime(row["date"], "%d/%m/%Y"))
    for index, row in enumerate(rows, start=1):
        row["draw_id"] = f"#{index:05d}"
    return rows


def parse_minhngoc_mega645_html(html_text):
    """Parse the Mega 6/45 block published on Minh Ngoc."""
    if not html_text:
        return None

    text = html_lib.unescape(html_text).replace("&nbsp;", " ").replace("\xa0", " ")
    draw_match = re.search(
        r"DT6X45_KY_VE[^>]*>\s*#?(\d+).*?"
        r"Ngày quay thưởng\s*(\d{1,2}/\d{1,2}/\d{4})",
        text,
        re.I | re.S,
    )
    if not draw_match:
        return None

    numbers_block = re.search(
        r'<ul[^>]*class=["\'][^"\']*result-number[^"\']*["\'][^>]*>'
        r"(.*?)(?:</ul>)",
        text,
        re.I | re.S,
    )
    if not numbers_block:
        return None

    nums = [
        int(value)
        for value in re.findall(
            r'class=["\'][^"\']*finnish[1-6][^"\']*["\'][^>]*>\s*(\d{1,2})\s*<',
            numbers_block.group(1),
            re.I,
        )
    ]
    if len(nums) != 6 or len(set(nums)) != 6 or not all(1 <= num <= 45 for num in nums):
        return None

    draw_date_str = draw_match.group(2)
    dt = datetime.datetime.strptime(draw_date_str, "%d/%m/%Y")
    weekdays = [
        "Thứ 2",
        "Thứ 3",
        "Thứ 4",
        "Thứ 5",
        "Thứ 6",
        "Thứ 7",
        "Chủ Nhật",
    ]
    jackpot_count_match = re.search(
        r'id=["\']DT6X45_S_JACKPOT["\'][^>]*>\s*([\d,.]+)',
        text,
        re.I,
    )
    jackpot_value_match = re.search(
        r'id=["\']DT6X45_G_JACKPOT["\'][^>]*>\s*([\d,.]+)',
        text,
        re.I,
    )

    def parse_amount(value):
        return int(value.replace(".", "").replace(",", "")) if value else 0

    return {
        "draw_id": f"#{int(draw_match.group(1)):05d}",
        "date": draw_date_str,
        "weekday": weekdays[dt.weekday()],
        "nums": sorted(nums),
        "jackpot": parse_amount(jackpot_value_match.group(1) if jackpot_value_match else ""),
        "winners": parse_amount(jackpot_count_match.group(1) if jackpot_count_match else ""),
    }


def add_manual_result(file_path, draw_id, date_str, nums, jackpot, winners=0):
    """Thêm dữ liệu kỳ quay theo cách thủ công (không cần API)."""
    # Lấy Thứ trong tuần
    dt = datetime.datetime.strptime(date_str, "%d/%m/%Y")
    weekdays = [
        "Thứ 2",
        "Thứ 3",
        "Thứ 4",
        "Thứ 5",
        "Thứ 6",
        "Thứ 7",
        "Chủ Nhật",
    ]
    weekday = weekdays[dt.weekday()]

    draw_data = {
        "draw_id": draw_id,
        "date": date_str,
        "weekday": weekday,
        "nums": sorted(nums),
        "jackpot": jackpot,
        "winners": winners,
    }
    
    append_to_excel(file_path, draw_data)


def fetch_latest_vietlott_mega645():
    """Lấy kết quả Mega 6/45 mới nhất. Hỗ trợ JSON cũ, HTML hiện tại và nhiều biến thể layout."""
    urls = [
        "https://vietlott.vn/vi/trung-thuong/ket-qua-trung-thuong/winning-number-645",
        "https://vietlott.vn/vi/trung-thuong/ket-qua-trung-thuong/645",
        "https://vietlott.vn/vi/trung-thuong/ket-qua-trung-thuong/mega-6-45",
        "https://www.minhngoc.net.vn/ket-qua-xo-so/dien-toan-vietlott.html",
        "https://vietlott.vn/api/front/get-result-mega645",
        "https://vietlott.vn/",
    ]
    headers = {
        "User-Agent": "Mozilla/5.0",
        "Accept-Language": "vi-VN,vi;q=0.9,en-US;q=0.8,en;q=0.7",
        "Referer": "https://vietlott.vn/",
    }

    for url in urls:
        try:
            print(f"Trying URL: {url}")
            res = requests.get(url, headers=headers, timeout=20)
            if hasattr(res, "raise_for_status"):
                res.raise_for_status()
            status_code = getattr(res, "status_code", 200)
            content_type = getattr(res, "headers", {}).get("Content-Type", "")
            print(f"Status Code: {status_code}")
            print(f"Content-Type: {content_type}")

            if "application/json" in content_type:
                data = res.json()
                latest = data.get("result", {})
                if not latest:
                    print("⚠️ API JSON trả về nhưng không có trường result.")
                    continue

                draw_id = f"#{int(latest['drawId']):05d}"
                draw_date_str = latest["drawDate"]
                nums = [int(n) for n in latest["winningNumbers"]]
                jackpot_val = int(latest["jackpotAmount"])
                jackpot_winners = int(latest["jackpotWinners"])

                dt = datetime.datetime.strptime(draw_date_str, "%d/%m/%Y")
                weekdays = [
                    "Thứ 2",
                    "Thứ 3",
                    "Thứ 4",
                    "Thứ 5",
                    "Thứ 6",
                    "Thứ 7",
                    "Chủ Nhật",
                ]
                weekday = weekdays[dt.weekday()]

                return {
                    "draw_id": draw_id,
                    "date": draw_date_str,
                    "weekday": weekday,
                    "nums": sorted(nums),
                    "jackpot": jackpot_val,
                    "winners": jackpot_winners,
                }

            minhngoc_parsed = parse_minhngoc_mega645_html(res.text)
            if minhngoc_parsed:
                print(
                    "✅ Parsed Minh Ngoc Mega 6/45: "
                    f"{minhngoc_parsed['draw_id']} {minhngoc_parsed['date']}"
                )
                return minhngoc_parsed

            parsed = parse_mega645_html(res.text)
            if parsed:
                print(f"✅ Parse HTML thành công: {parsed['draw_id']} {parsed['date']}")
                return parsed

            print("⚠️ HTML response does not contain a Mega 6/45 result block on this URL.")
        except Exception as e:
            print(f"Lỗi khi lấy dữ liệu Vietlott từ {url}: {e}")

    error_message = (
        "Không thể lấy được dữ liệu Mega 6/45 từ mọi URL fallback. "
        "The upstream sources may be blocking GitHub Actions or have changed."
    )
    print(f"❌ {error_message}")
    raise RuntimeError(error_message)


def init_sqlite_db(db_path="vietlott.db"):
    """Create the SQLite database and table used for lightweight storage."""
    db_file = Path(db_path)
    db_file.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(db_file)
    connection.execute(
        """
        CREATE TABLE IF NOT EXISTS results (
            draw_id TEXT PRIMARY KEY,
            draw_date TEXT NOT NULL,
            weekday TEXT NOT NULL,
            n1 INTEGER NOT NULL,
            n2 INTEGER NOT NULL,
            n3 INTEGER NOT NULL,
            n4 INTEGER NOT NULL,
            n5 INTEGER NOT NULL,
            n6 INTEGER NOT NULL,
            jackpot INTEGER DEFAULT 0,
            winners INTEGER DEFAULT 0,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
        """
    )
    connection.commit()
    connection.close()
    return str(db_file)


def fetch_ketquadientoan_history(date_from="20-07-2016", date_to="20-09-2026"):
    """Fetch the full Mega 6/45 archive from ketquadientoan.com."""
    url = (
        "https://www.ketquadientoan.com/tat-ca-ky-xo-so-mega-6-45.html"
        f"?datef={date_from}&datet={date_to}"
    )
    response = requests.get(
        url,
        timeout=30,
        headers={"User-Agent": "Mozilla/5.0"},
    )
    response.raise_for_status()
    return parse_ketquadientoan_archive_html(response.text)


def fetch_official_mega645_draw(draw_number):
    """Fetch one exact Mega 6/45 draw from Vietlott's official result page."""
    draw_number = int(draw_number)
    expected_id = f"#{draw_number:05d}"
    url = (
        "https://vietlott.vn/vi/trung-thuong/ket-qua-trung-thuong/645"
        f"?id={draw_number:05d}&nocatche=1"
    )
    response = requests.get(
        url,
        timeout=30,
        headers={"User-Agent": "Mozilla/5.0", "Accept-Language": "vi-VN,vi;q=0.9"},
    )
    response.raise_for_status()
    result = parse_mega645_html(response.text)
    if result is None or result["draw_id"] != expected_id:
        found = result["draw_id"] if result else "no result"
        raise ValueError(f"Expected {expected_id} at {url}, received {found}.")
    return result


def draw_id_to_number(draw_id):
    """Convert a draw id like '#01564' into its numeric value."""
    if not draw_id:
        return 0
    match = re.search(r"(\d+)", str(draw_id))
    return int(match.group(1)) if match else 0


def save_result_to_sqlite(db_path, draw_data):
    """Insert or update one Mega 6/45 result without removing older draws."""
    db_file = Path(db_path)
    db_file.parent.mkdir(parents=True, exist_ok=True)
    connection = sqlite3.connect(db_file)
    cursor = connection.cursor()
    cursor.execute(
        """
        CREATE TABLE IF NOT EXISTS results (
            draw_id TEXT PRIMARY KEY,
            draw_date TEXT NOT NULL,
            weekday TEXT NOT NULL,
            n1 INTEGER NOT NULL,
            n2 INTEGER NOT NULL,
            n3 INTEGER NOT NULL,
            n4 INTEGER NOT NULL,
            n5 INTEGER NOT NULL,
            n6 INTEGER NOT NULL,
            jackpot INTEGER DEFAULT 0,
            winners INTEGER DEFAULT 0,
            created_at TEXT DEFAULT CURRENT_TIMESTAMP
        )
        """
    )

    incoming_num = draw_id_to_number(draw_data.get("draw_id"))
    if incoming_num <= 0:
        connection.close()
        return False

    nums = list(draw_data["nums"])[:6]
    while len(nums) < 6:
        nums.append(0)

    cursor.execute(
        """
        INSERT INTO results (
            draw_id, draw_date, weekday, n1, n2, n3, n4, n5, n6, jackpot, winners
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        ON CONFLICT(draw_id) DO UPDATE SET
            draw_date = excluded.draw_date,
            weekday = excluded.weekday,
            n1 = excluded.n1,
            n2 = excluded.n2,
            n3 = excluded.n3,
            n4 = excluded.n4,
            n5 = excluded.n5,
            n6 = excluded.n6,
            jackpot = excluded.jackpot,
            winners = excluded.winners
        """,
        (
            draw_data["draw_id"],
            draw_data["date"],
            draw_data["weekday"],
            nums[0],
            nums[1],
            nums[2],
            nums[3],
            nums[4],
            nums[5],
            int(draw_data.get("jackpot", 0) or 0),
            int(draw_data.get("winners", 0) or 0),
        ),
    )
    connection.commit()
    connection.close()
    return True


def backfill_sqlite_history(db_path="vietlott.db", date_from="20-07-2016", date_to=None):
    """Fetch historical draws and store them in SQLite; existing records are updated."""
    if date_to is None:
        date_to = datetime.date.today().strftime("%d-%m-%Y")

    init_sqlite_db(db_path)
    history = fetch_ketquadientoan_history(date_from, date_to)
    latest = fetch_latest_vietlott_mega645()
    rows_by_draw = {row["draw_id"]: row for row in history}
    rows_by_draw[latest["draw_id"]] = latest

    saved = sum(save_result_to_sqlite(db_path, row) for row in rows_by_draw.values())
    print(f"SQLite backfill complete: {saved} Mega 6/45 draws stored in {db_path}.")
    return saved


def rebuild_sqlite_from_official_history(db_path="vietlott.db", workers=6):
    """Replace the SQLite result set with every official, ID-verified draw.

    Data is fetched and verified before the existing table is changed, avoiding
    a partially rebuilt database if the official site has a temporary error.
    """
    latest = fetch_latest_vietlott_mega645()
    last_draw = draw_id_to_number(latest["draw_id"])
    if last_draw <= 0:
        raise RuntimeError("Could not determine the latest official Mega 6/45 draw ID.")

    results = {}
    failures = []
    print(f"Fetching {last_draw} official Mega 6/45 draws with {workers} workers...")
    with ThreadPoolExecutor(max_workers=workers) as executor:
        futures = {
            executor.submit(fetch_official_mega645_draw, draw_number): draw_number
            for draw_number in range(1, last_draw + 1)
        }
        for completed, future in enumerate(as_completed(futures), start=1):
            draw_number = futures[future]
            try:
                result = future.result()
                results[result["draw_id"]] = result
            except Exception as error:
                failures.append(f"#{draw_number:05d}: {error}")
            if completed % 100 == 0 or completed == last_draw:
                print(f"Fetched {completed}/{last_draw} official draws.")

    if failures or len(results) != last_draw:
        details = "; ".join(failures[:10])
        raise RuntimeError(
            f"Official rebuild stopped without changing SQLite: {len(results)}/{last_draw} "
            f"draws verified. Failures: {details}"
        )

    init_sqlite_db(db_path)
    connection = sqlite3.connect(db_path)
    cursor = connection.cursor()
    cursor.execute("DELETE FROM results")
    cursor.executemany(
        """
        INSERT INTO results (
            draw_id, draw_date, weekday, n1, n2, n3, n4, n5, n6, jackpot, winners
        ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
        """,
        [
            (
                row["draw_id"], row["date"], row["weekday"], *row["nums"],
                int(row.get("jackpot", 0) or 0), int(row.get("winners", 0) or 0),
            )
            for row in sorted(results.values(), key=lambda item: draw_id_to_number(item["draw_id"]))
        ],
    )
    connection.commit()
    connection.close()
    print(f"Official SQLite rebuild complete: {last_draw} verified draws stored in {db_path}.")
    return last_draw


def render_results_html(db_path="vietlott.db", limit=20):
    """Return a simple HTML table with the latest stored Vietlott results."""
    connection = sqlite3.connect(db_path)
    connection.row_factory = sqlite3.Row
    rows = connection.execute(
        """
        SELECT draw_id, draw_date, weekday, n1, n2, n3, n4, n5, n6, jackpot, winners
        FROM results
        ORDER BY CAST(substr(draw_id, 2) AS INTEGER) DESC
        LIMIT ?
        """,
        (limit,),
    ).fetchall()
    connection.close()

    def number_balls(row):
        return "".join(
            f'<span class="number-ball">{int(row[f"n{index}"]):02d}</span>'
            for index in range(1, 7)
        )

    html_rows = []
    for row in rows:
        jackpot = int(row["jackpot"] or 0)
        html_rows.append(
            f"""
            <tr>
                <td class="draw-id">{html_lib.escape(str(row['draw_id']))}</td>
                <td class="date-cell">{html_lib.escape(str(row['draw_date']))}<span class="weekday">{html_lib.escape(str(row['weekday']))}</span></td>
                <td><div class="number-balls">{number_balls(row)}</div></td>
                <td class="jackpot-cell" data-raw-value="{jackpot}">{jackpot:,} đ</td>
                <td><span class="winner-count">{int(row['winners'] or 0)}</span></td>
            </tr>
            """.strip()
        )

    if rows:
        latest = rows[0]
        latest_jackpot = int(latest["jackpot"] or 0)
        latest_panel = f"""
            <section class="latest-draw" aria-labelledby="latest-heading">
                <div class="latest-meta">
                    <p class="eyebrow">LATEST DRAW <span>{html_lib.escape(str(latest['draw_id']))}</span></p>
                    <h2 id="latest-heading">{html_lib.escape(str(latest['draw_date']))}</h2>
                    <p class="latest-weekday">{html_lib.escape(str(latest['weekday']))}</p>
                </div>
                <div class="latest-balls" aria-label="Winning numbers">
                    {number_balls(latest)}
                </div>
                <div class="latest-prize">
                    <span class="eyebrow">JACKPOT</span>
                    <strong data-raw-value="{latest_jackpot}">{latest_jackpot:,} đ</strong>
                    <span>{int(latest['winners'] or 0)} jackpot winners</span>
                </div>
            </section>
        """
    else:
        latest_panel = '<p class="empty-state">No draw results are available yet.</p>'

    html_rows = "\n".join(html_rows)

    return f"""
    <!DOCTYPE html>
    <html lang="en">
    <head>
        <meta charset="utf-8" />
        <meta name="viewport" content="width=device-width, initial-scale=1" />
        <meta name="theme-color" content="#f5f7f2" />
        <meta name="description" content="Recent Vietlott Mega 6/45 draw results." />
        <title>Mega 6/45 Results</title>
        <link rel="preconnect" href="https://fonts.googleapis.com" />
        <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin />
        <link href="https://fonts.googleapis.com/css2?family=DM+Sans:wght@400;500;600;700&family=Space+Grotesk:wght@500;600;700&display=swap" rel="stylesheet" />
        <style>
            :root {{
                color-scheme: light;
                --ink: #172822;
                --muted: #68766f;
                --paper: #f5f7f2;
                --white: #ffffff;
                --green: #125d4e;
                --green-deep: #103b32;
                --line: #e2e8e1;
                --coral: #e45d46;
            }}
            * {{ box-sizing: border-box; }}
            body {{
                margin: 0;
                color: var(--ink);
                background-color: var(--paper);
                background-image: repeating-linear-gradient(135deg, transparent 0 22px, rgb(18 93 78 / 2.5%) 22px 23px);
                font-family: "DM Sans", sans-serif;
                font-size: 14px;
                -webkit-font-smoothing: antialiased;
            }}
            .site-header {{
                border-bottom: 1px solid var(--line);
                background: rgb(255 255 255 / 88%);
            }}
            .header-inner, main {{ width: min(1120px, calc(100% - 40px)); margin: 0 auto; }}
            .header-inner {{ min-height: 72px; display: flex; align-items: center; justify-content: space-between; gap: 20px; }}
            .brand {{ display: flex; align-items: center; gap: 12px; color: var(--ink); text-decoration: none; }}
            .brand-mark {{
                display: grid; place-items: center; width: 38px; height: 38px; border-radius: 50%;
                background: var(--coral); color: white; font: 700 12px "Space Grotesk", sans-serif;
            }}
            .brand-name {{ font: 700 16px "Space Grotesk", sans-serif; letter-spacing: 0; }}
            .header-note {{ color: var(--muted); font-size: 12px; text-align: right; }}
            main {{ padding: 42px 0 64px; }}
            .page-heading {{ display: flex; align-items: end; justify-content: space-between; gap: 20px; margin-bottom: 22px; }}
            .kicker, .eyebrow {{ margin: 0; color: #78c9a3; font-size: 10px; font-weight: 700; letter-spacing: 1.2px; }}
            .kicker {{ color: var(--green); }}
            h1 {{ margin: 7px 0 0; font: 700 clamp(26px, 4vw, 38px)/1.08 "Space Grotesk", sans-serif; letter-spacing: 0; }}
            .result-count {{ color: var(--muted); font-size: 12px; white-space: nowrap; }}
            .latest-draw {{
                display: grid; grid-template-columns: minmax(150px, .75fr) minmax(280px, 1.5fr) minmax(190px, .9fr);
                align-items: center; gap: 28px; padding: 26px 30px; border-radius: 8px;
                color: white; background: var(--green-deep);
                box-shadow: 0 12px 30px rgb(16 59 50 / 10%);
            }}
            .latest-meta h2 {{ margin: 10px 0 3px; font: 600 23px/1.15 "Space Grotesk", sans-serif; letter-spacing: 0; }}
            .latest-weekday, .latest-prize > span:last-child {{ margin: 0; color: #c2d5ce; font-size: 12px; }}
            .eyebrow span {{ margin-left: 7px; color: white; }}
            .latest-balls, .number-balls {{ display: flex; align-items: center; gap: 8px; }}
            .latest-balls {{ justify-content: center; flex-wrap: wrap; }}
            .number-ball {{
                display: inline-grid; place-items: center; flex: 0 0 36px; width: 36px; height: 36px;
                border: 1px solid #d9e9df; border-radius: 50%; background: #f7fbf6; color: var(--green-deep);
                font: 700 13px "Space Grotesk", sans-serif; font-variant-numeric: tabular-nums;
            }}
            .latest-balls .number-ball {{ flex-basis: 42px; width: 42px; height: 42px; border: 0; font-size: 15px; }}
            .latest-balls .number-ball:nth-child(3n + 2) {{ color: #b43e30; }}
            .latest-prize {{ display: flex; flex-direction: column; gap: 6px; padding-left: 22px; border-left: 1px solid rgb(255 255 255 / 18%); }}
            .latest-prize strong {{ font: 700 19px/1.2 "Space Grotesk", sans-serif; letter-spacing: 0; font-variant-numeric: tabular-nums; }}
            .history-heading {{ display: flex; align-items: baseline; justify-content: space-between; gap: 16px; margin: 38px 0 12px; }}
            .history-heading h2 {{ margin: 0; font: 600 19px "Space Grotesk", sans-serif; letter-spacing: 0; }}
            .history-heading span {{ color: var(--muted); font-size: 12px; }}
            .table-wrap {{ overflow-x: auto; border-top: 1px solid var(--line); border-bottom: 1px solid var(--line); background: var(--white); }}
            table {{ width: 100%; min-width: 760px; border-collapse: collapse; text-align: left; }}
            th {{ padding: 12px 14px; color: var(--muted); background: #edf3ed; font-size: 10px; font-weight: 700; letter-spacing: .7px; text-transform: uppercase; }}
            td {{ padding: 11px 14px; border-top: 1px solid #edf0eb; white-space: nowrap; }}
            tbody tr:hover {{ background: #f7faf6; }}
            .draw-id {{ color: var(--green); font-weight: 700; font-variant-numeric: tabular-nums; }}
            .date-cell {{ font-weight: 600; font-variant-numeric: tabular-nums; }}
            .weekday {{ display: block; margin-top: 2px; color: var(--muted); font-size: 11px; font-weight: 400; }}
            .jackpot-cell {{ font-weight: 600; font-variant-numeric: tabular-nums; }}
            .winner-count {{ display: inline-grid; place-items: center; min-width: 27px; height: 27px; border-radius: 50%; background: #edf3ed; color: var(--green); font-weight: 700; }}
            .empty-state {{ padding: 22px; color: var(--muted); background: white; }}
            .site-footer {{ width: min(1120px, calc(100% - 40px)); margin: 0 auto; padding: 16px 0 26px; border-top: 1px solid var(--line); color: var(--muted); font-size: 11px; }}
            @media (max-width: 760px) {{
                .header-inner, main, .site-footer {{ width: min(100% - 28px, 600px); }}
                .header-inner {{ min-height: 62px; }}
                .header-note {{ max-width: 130px; font-size: 11px; }}
                main {{ padding: 30px 0 44px; }}
                .page-heading {{ align-items: start; }}
                .latest-draw {{ grid-template-columns: 1fr; gap: 18px; padding: 22px; }}
                .latest-meta {{ display: grid; grid-template-columns: 1fr auto; align-items: baseline; gap: 4px 12px; }}
                .latest-meta .eyebrow {{ grid-column: 1 / -1; }}
                .latest-meta h2 {{ margin: 3px 0 0; }}
                .latest-weekday {{ text-align: right; }}
                .latest-balls {{ justify-content: flex-start; gap: 7px; }}
                .latest-prize {{ padding: 14px 0 0; border-left: 0; border-top: 1px solid rgb(255 255 255 / 18%); }}
                .history-heading {{ margin-top: 30px; }}
            }}
            @media (max-width: 380px) {{
                .number-ball {{ flex-basis: 32px; width: 32px; height: 32px; }}
                .latest-balls .number-ball {{ flex-basis: 38px; width: 38px; height: 38px; }}
                .latest-draw {{ padding: 18px; }}
            }}
        </style>
    </head>
    <body>
        <header class="site-header">
            <div class="header-inner">
                <a class="brand" href="./results.html" aria-label="Mega 6/45 results home">
                    <span class="brand-mark" aria-hidden="true">6/45</span>
                    <span class="brand-name">VIETLOTT RESULTS</span>
                </a>
                <span class="header-note">Official draw history<br />Latest results and jackpot details</span>
            </div>
        </header>
        <main>
            <div class="page-heading">
                <div><p class="kicker">MEGA 6/45 · DRAW ARCHIVE</p><h1>Winning numbers</h1></div>
                <span class="result-count">Showing {len(rows)} recent draws</span>
            </div>
            {latest_panel}
            <div class="history-heading"><h2>Recent history</h2><span>Newest draw first</span></div>
            <div class="table-wrap">
                <table>
                    <thead><tr><th scope="col">Draw</th><th scope="col">Date</th><th scope="col">Winning numbers</th><th scope="col">Jackpot</th><th scope="col">Winners</th></tr></thead>
                    <tbody>{html_rows}</tbody>
                </table>
            </div>
        </main>
        <footer class="site-footer">Historical results only. Lottery draws are random; past results do not predict future numbers.</footer>
    </body>
    </html>
    """


def append_to_excel(file_path, draw_data):
    """Ghi thêm dữ liệu kỳ mới vào file Excel."""
    wb = openpyxl.load_workbook(file_path)
    ws_data = wb["Lịch Sử Số Trúng (1200+ Kỳ)"]

    incoming_num = draw_id_to_number(draw_data.get("draw_id"))
    max_draw_num = 0
    for row_idx in range(1, ws_data.max_row + 1):
        cell_value = ws_data.cell(row=row_idx, column=1).value
        if cell_value:
            max_draw_num = max(max_draw_num, draw_id_to_number(cell_value))

    if incoming_num and max_draw_num and incoming_num < max_draw_num:
        print(
            f"Bỏ qua kỳ quay cũ {draw_data['draw_id']} vì Excel đã có kỳ mới hơn "
            f"({max_draw_num})."
        )
        return

    # Skip an exact duplicate, but repair stale rows where an old source
    # reused the same draw ID with a different date.
    last_row = ws_data.max_row
    last_draw_id = ws_data.cell(row=last_row, column=1).value

    if (
        last_draw_id == draw_data["draw_id"]
        and ws_data.cell(row=last_row, column=2).value == draw_data["date"]
    ):
        print(f"Kỳ quay {draw_data['draw_id']} đã tồn tại trong Excel!")
        return

    # Dùng thêm một lớp chống duplicate khi cùng draw_id xuất hiện ở nhiều nguồn hoặc chạy lại nhanh.
    for row_idx in range(1, last_row + 1):
        existing_draw_id = ws_data.cell(row=row_idx, column=1).value
        if existing_draw_id == draw_data["draw_id"]:
            existing_date = ws_data.cell(row=row_idx, column=2).value
            if existing_date != draw_data["date"]:
                print(
                    f"Kỳ quay {draw_data['draw_id']} có ngày cũ {existing_date}; "
                    f"cập nhật thành {draw_data['date']} ở hàng {row_idx}."
                )
                for col_idx, val in enumerate(
                    [
                        draw_data["draw_id"],
                        draw_data["date"],
                        draw_data["weekday"],
                        *draw_data["nums"],
                        draw_data["jackpot"],
                        draw_data["winners"],
                    ],
                    start=1,
                ):
                    cell = ws_data.cell(row=row_idx, column=col_idx, value=val)
                    if col_idx == 10:
                        cell.number_format = '#,##0 "VND"'
                        cell.alignment = Alignment(horizontal="right")
                wb.save(file_path)
                return
            print(f"Kỳ quay {draw_data['draw_id']} đã tồn tại ở hàng {row_idx}; bỏ qua ghi đè.")
            return

    # Chuẩn bị dòng mới
    new_row_idx = last_row + 1
    new_data = [
        draw_data["draw_id"],
        draw_data["date"],
        draw_data["weekday"],
        *draw_data["nums"],
        draw_data["jackpot"],
        draw_data["winners"],
    ]

    # Styles
    font_regular = Font(name="Calibri", size=11)
    font_num = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
    fill_gold = PatternFill(
        start_color="FEF3C7", end_color="FEF3C7", fill_type="solid"
    )
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    # Ghi dữ liệu & định dạng
    for col_idx, val in enumerate(new_data, start=1):
        cell = ws_data.cell(row=new_row_idx, column=col_idx, value=val)
        cell.border = thin_border
        cell.font = font_regular

        if col_idx in [1, 2, 3]:
            cell.alignment = Alignment(horizontal="center")
        elif 4 <= col_idx <= 9:
            cell.alignment = Alignment(horizontal="center")
            cell.number_format = "00"
            cell.font = font_num
        elif col_idx == 10:
            cell.alignment = Alignment(horizontal="right")
            cell.number_format = '#,##0 "VND"'
        elif col_idx == 11:
            cell.alignment = Alignment(horizontal="center")
            if val > 0:
                cell.fill = fill_gold
                cell.font = Font(
                    name="Calibri", size=11, bold=True, color="B45309"
                )

    wb.save(file_path)
    print(f"✅ Đã thêm thành công kỳ {draw_data['draw_id']} vào Excel!")


# Chạy bot
if __name__ == "__main__":
    if "--rebuild-official" in sys.argv:
        rebuild_sqlite_from_official_history()
        raise SystemExit(0)

    if "--backfill" in sys.argv:
        backfill_sqlite_history()
        raise SystemExit(0)

    result = fetch_latest_vietlott_mega645()
    append_to_excel("Vietlott_Mega_645_Full_Results.xlsx", result)
    init_sqlite_db("vietlott.db")
    save_result_to_sqlite("vietlott.db", result)
    with open("results.html", "w", encoding="utf-8") as file:
        file.write(render_results_html("vietlott.db", limit=20))
