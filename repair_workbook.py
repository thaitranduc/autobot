import re
import sqlite3

import openpyxl

from vietlott import fetch_ketquadientoan_history, fetch_latest_vietlott_mega645

SHEET_NAME = "Lịch Sử Số Trúng (1200+ Kỳ)"
HEADER = [
    "Kỳ Quay",
    "Ngày Quay",
    "Thứ",
    "Số 1",
    "Số 2",
    "Số 3",
    "Số 4",
    "Số 5",
    "Số 6",
    "Jackpot (VNĐ)",
    "Số Vé Trúng JP",
]


def draw_id_to_number(draw_id):
    if draw_id is None:
        return 0
    match = re.search(r"(\d+)", str(draw_id))
    return int(match.group(1)) if match else 0


def rebuild_workbook(file_path="Vietlott_Mega_645_Full_Results.xlsx"):
    archive_rows = fetch_ketquadientoan_history("20-07-2016", "20-09-2026")
    latest_row = fetch_latest_vietlott_mega645()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(HEADER)

    for row in archive_rows:
        ws.append([
            row["draw_id"],
            row["date"],
            row["weekday"],
            *row["nums"],
            row["jackpot"],
            row["winners"],
        ])

    last_draw_id = ws.cell(ws.max_row, 1).value if ws.max_row > 1 else None
    if last_draw_id != latest_row["draw_id"]:
        ws.append([
            latest_row["draw_id"],
            latest_row["date"],
            latest_row["weekday"],
            *latest_row["nums"],
            latest_row["jackpot"],
            latest_row["winners"],
        ])

    wb.save(file_path)
    print(f"Workbook rebuilt from archive page with {len(archive_rows)} archive rows + live row.")
    print(f"First draw: {ws.cell(2, 1).value} / {ws.cell(2, 2).value}")
    print(f"Last draw: {ws.cell(ws.max_row, 1).value} / {ws.cell(ws.max_row, 2).value}")


def rebuild_workbook_from_sqlite(
    file_path="Vietlott_Mega_645_Full_Results.xlsx",
    db_path="vietlott.db",
):
    """Rebuild the Excel history from the verified complete SQLite dataset."""
    with sqlite3.connect(db_path) as connection:
        rows = connection.execute(
            """
            SELECT draw_id, draw_date, weekday, n1, n2, n3, n4, n5, n6, jackpot, winners
            FROM results
            ORDER BY CAST(SUBSTR(draw_id, 2) AS INTEGER)
            """
        ).fetchall()

    if not rows:
        raise RuntimeError(f"No results found in {db_path}.")

    expected_ids = list(range(1, len(rows) + 1))
    actual_ids = [draw_id_to_number(row[0]) for row in rows]
    if actual_ids != expected_ids:
        raise RuntimeError("SQLite draw IDs are not a complete sequential history.")

    wb = openpyxl.load_workbook(file_path)
    ws = wb[SHEET_NAME]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for row in rows:
        ws.append(list(row))

    wb.save(file_path)
    print(f"Workbook rebuilt from SQLite with {len(rows)} complete draw rows.")
    print(f"First draw: {rows[0][0]} / {rows[0][1]}")
    print(f"Last draw: {rows[-1][0]} / {rows[-1][1]}")


if __name__ == "__main__":
    rebuild_workbook_from_sqlite()
