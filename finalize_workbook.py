import openpyxl
from vietlott import fetch_ketquadientoan_history, fetch_latest_vietlott_mega645

SHEET_NAME = "Lịch Sử Số Trúng (1200+ Kỳ)"
HEADER = [
    "Kỳ Quay",
    "Ngày Quay",
    "Thứ",
    "Số 1",
    "Số 2",
    "Số 3",
    "Số 4",
    "Số 5",
    "Số 6",
    "Jackpot (VNĐ)",
    "Số Vé Trúng JP",
]

archive_rows = fetch_ketquadientoan_history("20-07-2016", "20-09-2026")
latest_row = fetch_latest_vietlott_mega645()

wb = openpyxl.load_workbook("Vietlott_Mega_645_Full_Results.xlsx")
ws = wb[SHEET_NAME]

while ws.max_row > 0:
    ws.delete_rows(1, 1)

ws.append(HEADER)
for row in archive_rows:
    ws.append([
        row["draw_id"],
        row["date"],
        row["weekday"],
        *row["nums"],
        row["jackpot"],
        row["winners"],
    ])

if ws.max_row == 1 or ws.cell(ws.max_row, 1).value != latest_row["draw_id"]:
    ws.append([
        latest_row["draw_id"],
        latest_row["date"],
        latest_row["weekday"],
        *latest_row["nums"],
        latest_row["jackpot"],
        latest_row["winners"],
    ])

wb.save("Vietlott_Mega_645_Full_Results.xlsx")
print(f"rows={ws.max_row}")
print(f"first={[ws.cell(2, c).value for c in range(1, 12)]}")
print(f"last={[ws.cell(ws.max_row, c).value for c in range(1, 12)]}")
